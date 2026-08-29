#!/usr/bin/env python3
"""Render the three V1.0 filing manuals from their canonical Markdown sources.

The generated Word files use one explicit ``compact_reference_guide`` design
system with a named Chinese-font override.  Markdown remains the reviewable
source of truth; the DOCX files are deterministic delivery artifacts and are
written below ``dist/software-copyright/documents`` by default.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path
import re
from typing import Iterable

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "dist/software-copyright/documents"
SOURCES = (
    ("软件说明书", ROOT / "docs/software-copyright/软件说明书.md"),
    ("设计说明书", ROOT / "docs/software-copyright/设计说明书.md"),
    ("用户手册", ROOT / "docs/software-copyright/用户手册.md"),
)

INK = "0B2545"
HEADING = "2E74B5"
HEADING_DARK = "1F4D78"
MUTED = "5B6573"
TABLE_HEADER = "E8EEF5"
TABLE_BORDER = "AEB8C4"
CALLOUT_FILL = "F4F6F9"
CODE_FILL = "F2F4F7"
PLACEHOLDER_FILL = "FFF2CC"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = (80, 80, 120, 120)


@dataclass(frozen=True)
class Numbering:
    bullet_abstract_id: int
    decimal_abstract_id: int
    bullet_num_id: int


def _set_font(run, *, size: float | None = None, bold: bool | None = None,
              color: str | None = None, italic: bool | None = None,
              ascii_name: str = "Calibri", east_asia: str = "Microsoft YaHei") -> None:
    run.font.name = ascii_name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), ascii_name)
    fonts.set(qn("w:hAnsi"), ascii_name)
    fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def _set_style_font(style, *, size: float, color: str = "000000",
                    bold: bool = False, ascii_name: str = "Calibri") -> None:
    style.font.name = ascii_name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), ascii_name)
    fonts.set(qn("w:hAnsi"), ascii_name)
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _paragraph_shading(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shading = ppr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        ppr.append(shading)
    shading.set(qn("w:fill"), fill)


def _paragraph_border(paragraph, *, color: str = TABLE_BORDER) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "4")
        node.set(qn("w:color"), color)
        borders.append(node)


def _configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    _set_style_font(normal, size=11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_values = {
        "Heading 1": (16, HEADING, 18, 10),
        "Heading 2": (13, HEADING, 14, 7),
        "Heading 3": (12, HEADING_DARK, 10, 5),
    }
    for name, (size, color, before, after) in heading_values.items():
        style = styles[name]
        _set_style_font(style, size=size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.0

    for name, size, color, bold in (
        ("JYS Cover Kicker", 11, HEADING, True),
        ("JYS Cover Title", 28, INK, True),
        ("JYS Cover Subtitle", 15, HEADING_DARK, False),
        ("JYS Cover Meta", 10.5, MUTED, False),
        ("JYS Code Block", 9, INK, False),
        ("JYS Note", 10.5, INK, False),
        ("JYS Table", 9.5, "000000", False),
    ):
        if name not in styles:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[name]
        _set_style_font(
            style,
            size=size,
            color=color,
            bold=bold,
            ascii_name="Consolas" if name == "JYS Code Block" else "Calibri",
        )
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15 if name == "JYS Code Block" else 1.25


def _configure_section(document: Document, document_type: str) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(f"鉴源盾内容来源可信取证系统 V1.0  |  {document_type}")
    _set_font(run, size=8.5, color=MUTED)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    label = paragraph.add_run("V1.0  ·  第 ")
    _set_font(label, size=8.5, color=MUTED)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    field_begin_run = paragraph.add_run()
    field_begin_run._r.append(field_begin)
    instruction_run = paragraph.add_run()
    instruction_run._r.append(instruction)
    field_end_run = paragraph.add_run()
    field_end_run._r.append(field_end)
    suffix = paragraph.add_run(" 页")
    _set_font(suffix, size=8.5, color=MUTED)

    first_footer = section.first_page_footer
    paragraph = first_footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("软件著作权申请材料草案  ·  含待申请人确认字段")
    _set_font(run, size=8.5, color=MUTED)


def _next_numeric_id(elements: Iterable, attribute: str) -> int:
    values = []
    for element in elements:
        value = element.get(qn(attribute))
        if value is not None and value.isdigit():
            values.append(int(value))
    return max(values, default=0) + 1


def _abstract_numbering(numbering, *, kind: str) -> int:
    root = numbering.element
    abstract_id = _next_numeric_id(root.findall(qn("w:abstractNum")), "w:abstractNumId")
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level.append(level_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    level.append(suffix)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    ppr.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    level.append(ppr)
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rpr.append(fonts)
    level.append(rpr)
    abstract.append(level)
    root.insert(0, abstract)
    return abstract_id


def _number_instance(numbering, abstract_id: int) -> int:
    root = numbering.element
    num_id = _next_numeric_id(root.findall(qn("w:num")), "w:numId")
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    # Word may otherwise continue a previous list that shares the same
    # abstract definition, even though it has a distinct numId.
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    root.append(num)
    return num_id


def _configure_numbering(document: Document) -> Numbering:
    numbering = document.part.numbering_part
    bullet_abstract = _abstract_numbering(numbering, kind="bullet")
    decimal_abstract = _abstract_numbering(numbering, kind="decimal")
    return Numbering(
        bullet_abstract_id=bullet_abstract,
        decimal_abstract_id=decimal_abstract,
        bullet_num_id=_number_instance(numbering, bullet_abstract),
    )


def _apply_numbering(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


INLINE_RE = re.compile(
    r"(\[待[^\]]*\]|`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\(https?://[^)]+\))"
)


def _add_inline(paragraph, text: str, *, base_size: float | None = None,
                base_bold: bool = False) -> None:
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            _set_font(run, size=base_size, bold=base_bold)
        token = match.group(0)
        if token.startswith("[待"):
            run = paragraph.add_run(token)
            _set_font(run, size=base_size, bold=True, color="7A5A00")
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), PLACEHOLDER_FILL)
            run._r.get_or_add_rPr().append(shading)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            _set_font(
                run,
                size=(base_size - 0.5) if base_size else 10,
                bold=base_bold,
                color=HEADING_DARK,
                ascii_name="Consolas",
            )
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_font(run, size=base_size, bold=True)
        else:
            label, url = token[1:].split("](", 1)
            run = paragraph.add_run(f"{label}（{url[:-1]}）")
            _set_font(run, size=base_size, bold=base_bold, color=HEADING)
            run.font.underline = True
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        _set_font(run, size=base_size, bold=base_bold)


def _cover(document: Document, document_type: str) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(92)
    kicker = document.add_paragraph(style="JYS Cover Kicker")
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.add_run("JianYuanShield  ·  SOFTWARE COPYRIGHT MATERIAL")
    title = document.add_paragraph(style="JYS Cover Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title.add_run("鉴源盾内容来源可信取证系统")
    subtitle = document.add_paragraph(style="JYS Cover Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    subtitle.add_run("V1.0  " + document_type)
    meta = document.add_paragraph(style="JYS Cover Meta")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(4)
    _add_inline(meta, "编制/审核：[待申请人填写]", base_size=10.5)
    meta = document.add_paragraph(style="JYS Cover Meta")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_inline(meta, "定稿日期：[待申请人填写：YYYY 年 MM 月 DD 日]", base_size=10.5)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _set_repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    trpr.append(node)


def _cell_margins(cell) -> None:
    top, bottom, start, end = CELL_MARGINS_DXA
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for edge, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _table_widths(columns: int, first_header: str) -> list[int]:
    if columns == 1:
        return [CONTENT_WIDTH_DXA]
    if columns == 2:
        return [2700, 6660]
    if columns == 3:
        return [1800, 3780, 3780]
    if columns == 4 and first_header in {"序号", "编号", "#"}:
        return [800, 2200, 3180, 3180]
    if columns == 4:
        return [1800, 2520, 2520, 2520]
    base, remainder = divmod(CONTENT_WIDTH_DXA, columns)
    return [base + (1 if index < remainder else 0) for index in range(columns)]


def _set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    tblpr = table._tbl.tblPr
    width = tblpr.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        tblpr.append(width)
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    indent = tblpr.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tblpr.append(indent)
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")
    layout = tblpr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)

    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        for cell, value in zip(row.cells, widths):
            cell.width = Inches(value / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.first_child_found_in("w:tcW")
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(value))
            tcw.set(qn("w:type"), "dxa")
            _cell_margins(cell)


def _add_table(document: Document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    normalized = [row + [""] * (columns - len(row)) for row in rows]
    table = document.add_table(rows=len(normalized), cols=columns)
    table.style = "Table Grid"
    _set_repeat_header(table.rows[0])
    for row_index, values in enumerate(normalized):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            paragraph = cell.paragraphs[0]
            paragraph.style = "JYS Table"
            paragraph.paragraph_format.space_after = Pt(0)
            if row_index == 0 and len(normalized) > 1:
                # Keep the heading attached to at least one data row. Word's
                # repeating-header flag alone can still strand a header at the
                # bottom of a page before the first table body row.
                paragraph.paragraph_format.keep_with_next = True
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if row_index == 0 or (column_index == 0 and len(value) <= 12)
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            _add_inline(paragraph, value, base_size=9.5, base_bold=row_index == 0)
            if row_index == 0:
                tcpr = cell._tc.get_or_add_tcPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), TABLE_HEADER)
                tcpr.append(shading)
    _set_table_geometry(table, _table_widths(columns, normalized[0][0]))
    after = document.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(2)


def _split_table_row(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def _is_table_separator(line: str) -> bool:
    cells = _split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph(style="JYS Code Block")
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    _paragraph_shading(paragraph, CODE_FILL)
    _paragraph_border(paragraph)
    run = paragraph.add_run("\n".join(lines))
    _set_font(run, size=9, color=INK, ascii_name="Consolas")


def _add_note(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="JYS Note")
    paragraph.paragraph_format.left_indent = Inches(0.15)
    paragraph.paragraph_format.right_indent = Inches(0.15)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    _paragraph_shading(paragraph, CALLOUT_FILL)
    _paragraph_border(paragraph, color=HEADING)
    _add_inline(paragraph, text)


def _render_markdown(document: Document, source: str, numbering: Numbering) -> None:
    lines = source.replace("\r\n", "\n").split("\n")
    index = 0
    paragraph_lines: list[str] = []
    decimal_num_id: int | None = None
    previous_block = ""

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        if paragraph_lines:
            paragraph = document.add_paragraph()
            _add_inline(paragraph, " ".join(part.strip() for part in paragraph_lines))
            paragraph_lines = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if index == 0 and stripped.startswith("# "):
            index += 1
            continue
        if not stripped:
            flush_paragraph()
            decimal_num_id = None
            previous_block = "blank"
            index += 1
            continue
        if stripped.startswith("```"):
            flush_paragraph()
            code: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            if index >= len(lines):
                raise ValueError("unclosed Markdown code fence")
            _add_code_block(document, code)
            previous_block = "code"
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and _is_table_separator(lines[index + 1]):
            flush_paragraph()
            rows = [_split_table_row(line)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_split_table_row(lines[index]))
                index += 1
            _add_table(document, rows)
            previous_block = "table"
            decimal_num_id = None
            continue
        heading = re.match(r"^(#{2,6})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = min(len(heading.group(1)) - 1, 3)
            paragraph = document.add_paragraph(style=f"Heading {level}")
            _add_inline(paragraph, heading.group(2), base_bold=True)
            previous_block = "heading"
            decimal_num_id = None
            index += 1
            continue
        if re.fullmatch(r"-{3,}", stripped):
            flush_paragraph()
            previous_block = "rule"
            index += 1
            continue
        if stripped.startswith(">"):
            flush_paragraph()
            _add_note(document, stripped.lstrip("> "))
            previous_block = "note"
            decimal_num_id = None
            index += 1
            continue
        bullet = re.match(r"^\s*[-*+]\s+(.+)$", line)
        if bullet:
            flush_paragraph()
            value = bullet.group(1)
            if re.match(r"^\[[ xX]\]\s*", value):
                checked = value[1:2].lower() == "x"
                value = re.sub(r"^\[[ xX]\]\s*", "", value)
                paragraph = document.add_paragraph()
                marker = paragraph.add_run("☒ " if checked else "☐ ")
                _set_font(marker, size=11, color=HEADING_DARK)
                _add_inline(paragraph, value)
            else:
                paragraph = document.add_paragraph()
                _apply_numbering(paragraph, numbering.bullet_num_id)
                _add_inline(paragraph, value)
            previous_block = "bullet"
            decimal_num_id = None
            index += 1
            continue
        numbered = re.match(r"^\s*\d+\.\s+(.+)$", line)
        if numbered:
            flush_paragraph()
            if previous_block != "numbered" or decimal_num_id is None:
                decimal_num_id = _number_instance(
                    document.part.numbering_part,
                    numbering.decimal_abstract_id,
                )
            paragraph = document.add_paragraph()
            _apply_numbering(paragraph, decimal_num_id)
            _add_inline(paragraph, numbered.group(1))
            previous_block = "numbered"
            index += 1
            continue
        paragraph_lines.append(line)
        previous_block = "paragraph"
        decimal_num_id = None
        index += 1
    flush_paragraph()


def _build(document_type: str, source_path: Path, output_path: Path) -> None:
    source = source_path.read_text(encoding="utf-8-sig")
    if "鉴源盾内容来源可信取证系统 V1.0" not in source:
        raise ValueError(f"canonical software name missing from {source_path}")
    document = Document()
    document.core_properties.title = f"鉴源盾内容来源可信取证系统 V1.0 {document_type}"
    document.core_properties.subject = "软件著作权申请材料"
    document.core_properties.creator = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.keywords = ""
    _configure_styles(document)
    _configure_section(document, document_type)
    numbering = _configure_numbering(document)
    _cover(document, document_type)
    _render_markdown(document, source, numbering)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)
    output_dir = args.output_dir.expanduser().resolve()
    for document_type, source_path in SOURCES:
        output = output_dir / f"鉴源盾内容来源可信取证系统-V1.0-{document_type}.docx"
        _build(document_type, source_path, output)
        print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
